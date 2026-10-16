from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import subprocess,re,json

root=Path(__file__).resolve().parents[2]
sha='386b6e03538ced921641997f67e2cd38b5d3d322'
html=subprocess.check_output(['git','show',sha+':index.html'],cwd=root).decode('utf-8')
css=subprocess.check_output(['git','show',sha+':estilos.css'],cwd=root).decode('utf-8')
imgs=re.findall(r'<img\b[^>]*>',html)
ids=set(re.findall(r'\bid="([^"]+)"',html))
frags=re.findall(r'href="#([^"]+)"',html)
broken=[x for x in frags if x not in ids]
tree=subprocess.check_output(['git','ls-tree','-rl',sha],cwd=root).decode()
sizes={line.split('\t')[1]:int(line.split('\t')[0].split()[-1]) for line in tree.splitlines()}
imagebytes=sum(sizes[re.search(r'src="([^"]+)"',tag)[1]] for tag in imgs)
assert len(imgs)==11 and len(broken)==7
(root/'tmp/seo/evidence.json').write_text(json.dumps({'commit':sha,'images':len(imgs),'image_bytes':imagebytes,'broken_fragments':broken,'sizes':sizes},ensure_ascii=False,indent=2),encoding='utf-8')
d=Document(); sec=d.sections[0]
sec.page_height=Cm(29.7);sec.page_width=Cm(21)
sec.top_margin=sec.bottom_margin=Cm(2);sec.left_margin=sec.right_margin=Cm(2.2)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 s=d.styles[name];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(11)
d.styles['Normal'].paragraph_format.space_after=Pt(7)
d.styles['Normal'].paragraph_format.line_spacing=1.1
d.styles['Title'].font.size=Pt(25)
d.styles['Heading 1'].font.size=Pt(18)
d.styles['Heading 2'].font.size=Pt(12)
def p(t):d.add_paragraph(t)
def h(t):d.add_heading(t,2)
def page(t):d.add_page_break();d.add_heading(t,1)
def table(headers,rows,widths):
 t=d.add_table(rows=1,cols=len(headers));t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Cm(w)
 for c,txt in zip(t.rows[0].cells,headers):c.text=txt
 for row in rows:
  for c,txt in zip(t.add_row().cells,row):c.text=str(txt)
 for i,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr()
  if i==0:trpr.append(OxmlElement('w:tblHeader'))
  for c in row.cells:
   pr=c._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   pr.append(borders);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE6F1' if i==0 else ('F4F6F8' if i%2==0 else 'FFFFFF'));pr.append(sh)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(5);para.paragraph_format.space_before=Pt(5)
    for run in para.runs:run.font.size=Pt(10);run.bold=i==0
footer=sec.footer.paragraphs[0];footer.alignment=2
footer.add_run('Ambienta  |  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
d.add_paragraph('Informe de auditoría SEO técnico de Ambienta', 'Title')
d.add_paragraph('Rendimiento velocidad y estabilidad visual antes de las correcciones', 'Subtitle')
p('Fecha de elaboración 11 de octubre de 2026\nProyecto FINISHWEBAMBIENTA\nRepositorio https://github.com/raquel-castro/FINISHWEBAMBIENTA')
h('Resumen ejecutivo')
p('La versión previa a las correcciones presenta deficiencias verificables en navegación y preparación técnica para buscadores. Se identificaron siete enlaces a fragmentos inexistentes, ausencia de descripción y URL canónica, once imágenes sin atributos de dimensiones ni estrategia explícita de carga y una configuración de carrusel que puede desplazar el contenido fuera del área visible. El peso de las imágenes referenciadas asciende a '+f'{imagebytes:,}'.replace(',',' ')+' bytes, aproximadamente '+f'{imagebytes/1e6:.2f}'.replace('.',',')+' MB decimales.')
p('La prioridad es recuperar una navegación funcional y revisar el carrusel, después optimizar las imágenes y completar las señales de búsqueda. El código ya dispone de idioma español, título descriptivo, un encabezado principal, textos alternativos y adaptación básica a pantallas móviles.')
h('Objetivo y alcance')
p('Evaluar el estado de Ambienta antes de los cambios del 11 de octubre de 2026 y explicar dónde se realizan las auditorías, qué herramientas las ejecutan y qué resultados pueden demostrarse. El análisis comprende HTML, CSS, imágenes y estructura de archivos de la versión histórica.')
p('Base técnica '+sha+'\nFecha del commit 13 de agosto de 2025 a las 16:38:01 UTC menos 5. Se excluyen los 24 commits del 11 de octubre de 2026 registrados en la copia local revisada. Su contenido no se utiliza como evidencia de los problemas originales.')
p('Resultado general: existen incidencias técnicas confirmadas y riesgos de rendimiento que requieren pruebas. La revisión de código no permite declarar que Core Web Vitals aprueba o desaprueba, ni asignar una puntuación histórica de PageSpeed o Lighthouse.')
page('Metodología y trazabilidad de la auditoría')
h('Selección de la versión anterior')
p('Se examinó el historial Git y se tomó el último commit anterior al inicio del 11 de octubre de 2026 en la zona America/Lima. Se consultaron directamente sus objetos mediante git show y git ls-tree, sin revertir la rama ni alterar el sitio. Esta selección conserva el estado previo a todos los commits de ese día, incluidos los cambios funcionales.')
h('Procedimiento aplicado')
p('Se revisó la cabecera HTML para identificar título, descripción, canonical e idioma. Se compararon los destinos href con los identificadores id; se inventariaron imágenes y atributos; se sumaron los tamaños de los archivos referenciados; y se inspeccionaron reglas CSS de tamaños, tipografía y animaciones. Los resultados son reproducibles sobre el commit indicado.')
h('Cómo se interpretan los resultados')
table(['Tipo de evidencia','Interpretación'],[
('Confirmado en código','Una propiedad o ausencia observable en la versión histórica, con archivo y línea de referencia.'),
('Riesgo inferido','Una posible consecuencia de esa configuración. Requiere validar en navegador o medir rendimiento.'),
('No medido','Dato que necesita ejecución, tráfico real, servidor o informe histórico. No se asigna un valor supuesto.')],[4,12.6])
h('Límites de la reconstrucción histórica')
p('Este documento presenta una auditoría retrospectiva del código, no una recuperación de un informe externo ya ejecutado. No consta una medición histórica conservada de PageSpeed, GTmetrix o Search Console. Tampoco se prueba la configuración histórica del alojamiento, su caché, compresión, TLS o tiempos de respuesta.')
p('La dirección del repositorio identifica el código fuente; no es la URL de la tienda que debe ingresarse en PageSpeed. Una medición actual de la web publicada puede incluir las correcciones y no demuestra cómo rendía antes. Para repetir pruebas sobre la versión original debe servirse ese commit en un entorno separado y registrar que se trata de una reconstrucción.')
p('Evidencia primaria: index.html y estilos.css del commit de referencia; inventario git ls-tree. Las rutas y líneas citadas en este informe corresponden a esa versión, no a los archivos corregidos actuales.')
page('Dónde se realizan las auditorías y qué páginas las hacen')
p('Las auditorías se realizan en servicios web que analizan una URL pública, en herramientas del navegador que prueban una página abierta y en rastreadores instalados que recorren el sitio. Cada herramienta cubre una parte distinta; sus informes deben conservarse con fecha, URL y configuración.')
table(['Herramienta y acceso','Qué revisa y qué entrega'],[
('PageSpeed Insights\npagespeed.web.dev','Pruebas móvil y escritorio. Diagnóstico Lighthouse y datos de usuarios de CrUX cuando existen. Entrega métricas y oportunidades de mejora. [2]'),
('Lighthouse\nChrome DevTools','Auditoría en navegador o CLI de rendimiento, SEO y accesibilidad. Puede ejecutarse sobre la versión histórica servida localmente. [3]'),
('Google Search Console\nsearch.google.com/search-console','Indexación, inspección de URLs y Core Web Vitals de usuarios reales. Requiere verificar la propiedad del sitio. [1, 4]'),
('GTmetrix\ngtmetrix.com','Diagnóstico de rendimiento basado en Lighthouse, con informe y análisis de carga. Registrar la configuración de cada prueba. [5]'),
('Screaming Frog SEO Spider\nscreamingfrog.co.uk/seo-spider','Aplicación de rastreo: enlaces, respuestas HTTP, títulos, descripciones y canonical. No es una medición de usuarios reales. [6]')],[5.7,10.9])
h('Uso recomendado para Ambienta')
p('Empezar con Lighthouse sobre la versión histórica aislada. Complementar con un rastreo de navegación y recursos. Si esa reconstrucción se dispone en una URL de prueba accesible, utilizar PageSpeed y GTmetrix; evitar interpretar sus resultados como una medición tomada en 2025. Search Console se utiliza sobre la propiedad real y permite verificar indexación y experiencia de usuarios, cuando hay datos suficientes.')
h('Evidencia que debe guardarse')
p('Conservar el informe HTML o JSON, captura del resumen, URL probada, commit, fecha y hora, versión del navegador y herramienta, dispositivo, red, caché y número de ejecuciones. Registrar por separado móvil y escritorio. No comparar pruebas hechas con configuraciones distintas como si fueran equivalentes.')
page('Resultado del SEO técnico antes de las correcciones')
table(['Control','Resultado histórico','Evidencia'],[
('Título e idioma','Presentes','index.html 2 y 6'),
('Meta descripción','Ausente','head líneas 3 a 9'),
('URL canónica','Ausente','head líneas 3 a 9'),
('H1 principal','Presente y descriptivo','index.html 62'),
('Sitemap y robots.txt','No incluidos en el árbol Git','Inventario de archivos'),
('Enlaces a fragmentos','7 destinos inexistentes','index.html 32 a 100'),
('Textos alternativos','11 de 11 imágenes con alt','Etiquetas img'),
('Datos estructurados','No declarados','Sin JSON LD o microdatos')],[4.1,7.1,5.4])
h('Descripción y canonical')
p('La ausencia de meta description limita el control editorial del resumen que puede mostrar Google; no bloquea por sí sola la indexación. Google también genera snippets desde el contenido. La falta de canonical deja sin declarar la URL preferida, pero no demuestra que haya duplicados: debe comprobarse si la raíz e index.html u otras variantes sirven el mismo contenido. [7, 8]')
h('Rastreo e indexación')
p('No hay sitemap.xml ni robots.txt en la versión histórica del repositorio. En una web de una página, un sitemap es una mejora de descubrimiento de prioridad menor y no un requisito de indexación. La ausencia de robots.txt tampoco significa bloqueo. Para determinar el estado real hacen falta respuestas HTTP e inspección de URL en la propiedad publicada. [4, 9]')
h('Navegación y contenido')
p('Los destinos inexistentes son #contacto, #carrito, #cuenta, #proyectos, #cocina, #dormitorio y #escritorios. Son fallos de navegación dentro del documento; no equivalen a siete errores HTTP 404. Además, cuatro enlaces usan únicamente #, incluido Ver Todo el Catálogo y las tres redes sociales, y no proporcionan un destino útil.')
p('Los cinco alt del carrusel son Mueble 1 a Mueble 5: existen, pero describen poco el contenido. Conviene identificar el tipo de mueble o espacio mostrado. Los productos carecen de páginas propias; si el catálogo crece, páginas con información diferenciada pueden ampliar la cobertura de búsqueda. Los datos estructurados son una oportunidad condicionada a información comercial auténtica, no una garantía de posicionamiento.')
page('Core Web Vitals rendimiento velocidad y estabilidad visual')
p('Core Web Vitals evalúa carga, respuesta a interacción y estabilidad visual. La calificación de campo exige que las tres métricas cumplan sus umbrales en el percentil 75, separando móvil y escritorio. Una puntuación general de rendimiento no reemplaza este criterio. [1]')
table(['Métrica','Objetivo bueno','Resultado previo demostrable'],[
('LCP','Hasta 2,5 segundos','No medido. Hay riesgos de carga por imágenes y recursos CSS externos.'),
('INP','Hasta 200 milisegundos','No medido. La versión no incluye JavaScript propio; ello no acredita buena interacción.'),
('CLS','Hasta 0,1','No medido. El CSS ya reserva altura en varios bloques; revisar fuente y logo.')],[2.3,4.6,9.7])
h('Carga y posible impacto en LCP')
p('Las once imágenes no declaran loading, fetchpriority, srcset ni sizes. Sin carga diferida explícita, las imágenes inferiores pueden competir por ancho de banda con la portada. El candidato LCP puede ser una imagen o el título; solo una traza permite identificarlo. La primera imagen no tiene prioridad alta declarada y las cinco imágenes del carrusel suman 2 259 464 bytes. No debe aplicarse lazy a la imagen que resulte ser el LCP.')
p('La hoja estilos.css y la hoja remota de Google Fonts se enlazan desde head. No se declaran preconnect. La fuente ya solicita display=swap; este aspecto favorable no debe presentarse como ausente. Conviene medir si la conexión y sustitución de Roboto afectan la carga y el diseño antes de optimizar.')
h('Estabilidad visual y animaciones')
p('Ninguna imagen declara width y height en HTML, pero esto no demuestra automáticamente un CLS elevado. El hero tiene altura de 70vh, las imágenes de categorías 200px y las de producto 250px. Estas reglas reservan espacio. El logo tiene altura fija y ancho automático, y la sustitución de fuente puede cambiar medidas; ambos requieren observarse en una traza de carga.')
p('El contenedor de diapositivas mide 500% del área visible y la animación lo desplaza hasta translateX(-400%). Los porcentajes de transform se calculan respecto al propio elemento: -100% ya equivale a cinco anchos del área visible. Se infiere un riesgo alto de mostrar fondo vacío. Es un problema visual y funcional; el desplazamiento mediante transform no implica por sí solo aumento de CLS.')
h('Respuesta a interacción')
p('Lighthouse de navegación utiliza TBT como diagnóstico y no mide directamente INP sin interacciones. Debe probarse abrir el menú y usar los controles, además de recopilar datos de usuarios reales. Un botón sin comportamiento puede parecer rápido y aun así no cumplir su función. [1, 3]')
page('Inventario de imágenes y hallazgos priorizados')
p('Los tamaños corresponden a los blobs originales de Git. No representan bytes transferidos en una sesión, duración de descarga ni el peso total con fuentes, cabeceras y caché. MB se expresa en base decimal.')
table(['Imagen','Bytes originales'],[(k.replace('galeria/',''),f'{v:,}'.replace(',',' ')) for k,v in sizes.items() if k.startswith('galeria/')]+[('Total imágenes',f'{imagebytes:,}'.replace(',',' '))],[11.5,5.1])
h('Prioridad alta')
p('H1. Corregir los siete destinos inexistentes y reemplazar los enlaces # por destinos funcionales. Verificación: cada enlace lleva al bloque o acción prevista y no deja al usuario sin respuesta. Evidencia: index.html 32, 36, 40, 65, 91, 95 y 99.')
p('H2. Revisar las proporciones y desplazamientos del carrusel. Verificación: mostrar todas las diapositivas durante un ciclo completo de 50 segundos sin áreas vacías, en móvil y escritorio. Evidencia: estilos.css 167 y bloque @keyframes slideShow.')
p('H3. Optimizar recursos de imagen y priorizar la portada. Emplear dimensiones apropiadas, formatos eficientes y carga diferida para contenido inferior. Verificación: comparar bytes y LCP con la misma configuración; conservar calidad visual.')
h('Prioridad media y menor')
p('H4. Incorporar descripción y canonical conforme a la URL real; revisar la reserva de espacio, Roboto y conexiones a fuentes con mediciones. H5. Mejorar textos alternativos y nombres accesibles de controles. H6. Incorporar sitemap si ayuda a administrar el sitio; revisar datos estructurados únicamente para contenido y negocio verificables.')
page('Plan de validación y conclusión de la auditoría')
h('Pruebas necesarias para completar el resultado cuantitativo')
p('1. Extraer el commit de referencia en una carpeta aislada y servirlo por HTTP. No reemplazar el sitio de producción ni mezclar archivos modernos. Revisar recursos ausentes, navegación, formulario, botones y carrusel en anchos de móvil y escritorio.')
p('2. Ejecutar Lighthouse al menos tres veces por perfil con caché y condiciones consistentes. Guardar el informe y resumir la mediana de LCP, CLS, TBT y puntuaciones. Esas cifras serían una medición nueva de una reconstrucción histórica, no resultados capturados antes de los commits.')
p('3. Probar interacciones con Chrome DevTools y confirmar el elemento LCP y las causas de cambios de diseño. Medir INP con interacciones o datos reales; no convertir TBT en un supuesto INP.')
p('4. Consultar Search Console y CrUX de la propiedad real cuando existan datos. Anotar el periodo observado: el agregado de campo puede abarcar versiones anteriores y posteriores. Verificar respuestas HTTP, indexación, canonical elegida y recursos del servidor.')
h('Criterios de aceptación')
p('Todos los destinos de navegación deben existir y realizar su función; el carrusel debe mantenerse visible; no deben faltar imágenes ni hojas de estilo. Para experiencia de campo, buscar LCP ≤ 2,5 s, INP ≤ 200 ms y CLS ≤ 0,1 en el percentil 75 por dispositivo. Las mejoras SEO deben validarse en HTML y después en la propiedad publicada. [1, 4]')
h('Conclusión')
p('Antes de los commits excluidos, Ambienta contaba con una estructura HTML básica favorable, pero presentaba fallos concretos de navegación y oportunidades claras de optimización. El mayor volumen de recursos está en las imágenes, especialmente el carrusel; la configuración de desplazamiento requiere atención inmediata. La ausencia de descripción y canonical reduce la preparación técnica, sin demostrar por sí sola pérdida de posiciones o falta de indexación.')
p('El resultado acreditado es una auditoría estática con incidencias confirmadas. El estado numérico de Core Web Vitals y las puntuaciones de herramientas externas permanecen sin medir. No hay fundamento para asignarles valores ni atribuir a una plataforma una auditoría que no consta ejecutada.')
page('Referencias y evidencia reproducible')
p('Fuentes oficiales consultadas el 11 de octubre de 2026. Los números entre corchetes remiten a las referencias siguientes.')
refs=[
('[1] Google web.dev Web Vitals','https://web.dev/articles/vitals'),
('[2] Google About PageSpeed Insights','https://developers.google.com/speed/docs/insights/v5/about'),
('[3] Chrome Introduction to Lighthouse','https://developer.chrome.com/docs/lighthouse/overview'),
('[4] Google Debugging your pages y Search Console','https://developers.google.com/search/help/debug'),
('[5] GTmetrix Informes basados en Lighthouse','https://gtmetrix.com/blog/welcome-to-the-new-gtmetrix-powered-by-lighthouse/'),
('[6] Screaming Frog SEO Spider','https://www.screamingfrog.co.uk/seo-spider/'),
('[7] Google Meta descriptions y snippets','https://developers.google.com/search/docs/appearance/snippet'),
('[8] Google Canonicalización','https://developers.google.com/search/docs/crawling-indexing/consolidate-duplicate-urls'),
('[9] Google Sitemaps','https://developers.google.com/search/docs/crawling-indexing/sitemaps/overview')]
for title,url in refs:p(title+'\n'+url)
h('Código original revisado')
p('HTML https://github.com/raquel-castro/FINISHWEBAMBIENTA/blob/'+sha+'/index.html\nCSS https://github.com/raquel-castro/FINISHWEBAMBIENTA/blob/'+sha+'/estilos.css')
h('Comandos para reproducir la inspección')
p('git show '+sha+':index.html\ngit show '+sha+':estilos.css\ngit ls-tree -rl '+sha)
d.core_properties.title='Informe de auditoría SEO técnico de Ambienta'
d.core_properties.subject='Versión previa a las correcciones del 11 de octubre de 2026'
d.save(root/'output/Informe_auditoria_SEO_Ambienta_antes_correcciones.docx')
print('Created',imagebytes,'bytes; broken anchors',broken)
